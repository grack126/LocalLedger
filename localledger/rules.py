"""Business rules: export templates, description builder, credit-note detection.

Sections
--------
1. Template engine  — ``TemplateConfig``, built-in presets, ``map_rows`` -> CSV/XLSX
2. Service period   — regex detection of billing periods in document text
3. Description rules — token templates such as ``"{Service_Month} - {Raw_Description}"``
4. Credit notes      — keyword / negative-total detection and PI->PC / SI->SC flipping
5. Archive naming    — ``[AccountCode]_[InvoiceDate]_[InvoiceNumber].pdf`` patterns
6. ``apply_rules``   — convenience pipeline used after extraction

Nothing in this module touches the database, so every function is unit-testable.
"""
from __future__ import annotations

import copy
import csv
import io
import json
import re
from dataclasses import asdict, dataclass, field
from datetime import date, timedelta
from decimal import ROUND_HALF_UP, Decimal
from typing import Any, Iterable, Mapping, Sequence

from .models import (
    ALL_FIELDS,
    AMOUNT_FIELDS,
    CENT,
    DocType,
    F,
    StagingRow,
    clean_str,
    month_label,
    parse_date,
)

# =============================================================================
# 1. TEMPLATE ENGINE
# =============================================================================
DATE_FORMATS: dict[str, str] = {
    "DD/MM/YYYY": "%d/%m/%Y",
    "YYYY-MM-DD": "%Y-%m-%d",
    "MM/DD/YYYY": "%m/%d/%Y",
    "DD-MM-YYYY": "%d-%m-%Y",
    "DD.MM.YYYY": "%d.%m.%Y",
    "DD/MM/YY": "%d/%m/%y",
    "YYYYMMDD": "%Y%m%d",
}
FILE_FORMATS = ("csv", "xlsx")
AMOUNT_SIGNS = ("positive", "negative_for_credits")
SERVICE_PERIOD_STYLES = ("month", "range")
ALL_DOC_TYPES = [t.value for t in DocType]


@dataclass
class ColumnSpec:
    """One output column.

    ``source`` is an internal field name (see ``models.ALL_FIELDS``) or ``None``
    for a static column. ``default`` fills static columns and blank values.
    """

    header: str
    source: str | None = None
    default: str = ""
    max_length: int | None = None

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> ColumnSpec:
        max_length = data.get("max_length")
        return cls(
            header=str(data["header"]),
            source=data.get("source") or None,
            default=clean_str(data.get("default", "")),
            max_length=int(max_length) if max_length not in (None, "", 0) else None,
        )


@dataclass
class TemplateConfig:
    name: str
    columns: list[ColumnSpec]
    software: str = "Custom"
    date_format: str = "DD/MM/YYYY"
    file_format: str = "csv"
    delimiter: str = ","
    encoding: str = "utf-8"
    include_header: bool = True
    decimal_places: int = 2
    amount_sign: str = "positive"
    applies_to: list[str] = field(default_factory=lambda: list(ALL_DOC_TYPES))
    service_period_style: str = "month"
    # field -> {lookup key -> output value}. Keys may be qualified as
    # "PI:T1" (doc type), "purchase:T1"/"sales:T1" (ledger) or plain "T1".
    value_maps: dict[str, dict[str, str]] = field(default_factory=dict)
    notes: str = ""
    # persistence metadata (not part of the JSON config)
    id: int | None = None
    is_builtin: bool = False
    builtin_version: int = 0

    # ------------------------------------------------------------- helpers
    @property
    def strftime(self) -> str:
        return DATE_FORMATS.get(self.date_format, self.date_format)

    @property
    def file_extension(self) -> str:
        return ".xlsx" if self.file_format == "xlsx" else ".csv"

    @property
    def mime_type(self) -> str:
        if self.file_format == "xlsx":
            return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        return "text/csv"

    def validate(self) -> list[str]:
        errors: list[str] = []
        if not self.name.strip():
            errors.append("Template name is required.")
        if not self.columns:
            errors.append("Template needs at least one column.")
        headers = [c.header.strip() for c in self.columns]
        if any(not h for h in headers):
            errors.append("Every column needs a header.")
        dupes = sorted({h for h in headers if h and headers.count(h) > 1})
        if dupes:
            errors.append(f"Duplicate column headers: {', '.join(dupes)}")
        for col in self.columns:
            if col.source and col.source not in ALL_FIELDS:
                errors.append(f"Column '{col.header}' uses unknown field '{col.source}'.")
            if col.max_length is not None and col.max_length < 1:
                errors.append(f"Column '{col.header}' has an invalid max length.")
        if self.date_format not in DATE_FORMATS and "%" not in self.date_format:
            errors.append(f"Unsupported date format '{self.date_format}'.")
        if self.file_format not in FILE_FORMATS:
            errors.append(f"File format must be one of {FILE_FORMATS}.")
        if self.file_format == "csv" and len(self.delimiter) != 1:
            errors.append("CSV delimiter must be a single character.")
        if not 0 <= self.decimal_places <= 6:
            errors.append("Decimal places must be between 0 and 6.")
        if self.amount_sign not in AMOUNT_SIGNS:
            errors.append(f"Amount sign must be one of {AMOUNT_SIGNS}.")
        if self.service_period_style not in SERVICE_PERIOD_STYLES:
            errors.append(f"Service period style must be one of {SERVICE_PERIOD_STYLES}.")
        bad_types = [t for t in self.applies_to if t not in ALL_DOC_TYPES]
        if bad_types or not self.applies_to:
            errors.append("Template must apply to at least one of PI, PC, SI, SC.")
        try:
            "x".encode(self.encoding)
        except LookupError:
            errors.append(f"Unknown text encoding '{self.encoding}'.")
        return errors

    def to_dict(self) -> dict[str, Any]:
        """Config only — id / builtin flags live in their own DB columns."""
        return {
            "name": self.name, "software": self.software,
            "date_format": self.date_format, "file_format": self.file_format,
            "delimiter": self.delimiter, "encoding": self.encoding,
            "include_header": self.include_header, "decimal_places": self.decimal_places,
            "amount_sign": self.amount_sign, "applies_to": list(self.applies_to),
            "service_period_style": self.service_period_style,
            "value_maps": copy.deepcopy(self.value_maps), "notes": self.notes,
            "columns": [c.to_dict() for c in self.columns],
        }

    def to_json(self) -> str:
        return json.dumps(self.to_dict(), ensure_ascii=False)

    @classmethod
    def from_dict(cls, data: Mapping[str, Any], **meta: Any) -> TemplateConfig:
        return cls(
            name=str(data["name"]),
            columns=[c if isinstance(c, ColumnSpec) else ColumnSpec.from_dict(c)
                     for c in data.get("columns", [])],
            software=data.get("software", "Custom"),
            date_format=data.get("date_format", "DD/MM/YYYY"),
            file_format=data.get("file_format", "csv"),
            delimiter=data.get("delimiter", ","),
            encoding=data.get("encoding", "utf-8"),
            include_header=bool(data.get("include_header", True)),
            decimal_places=int(data.get("decimal_places", 2)),
            amount_sign=data.get("amount_sign", "positive"),
            applies_to=list(data.get("applies_to") or ALL_DOC_TYPES),
            service_period_style=data.get("service_period_style", "month"),
            value_maps={k: dict(v) for k, v in (data.get("value_maps") or {}).items()},
            notes=data.get("notes", ""),
            **meta,
        )

    @classmethod
    def from_json(cls, text: str, **meta: Any) -> TemplateConfig:
        return cls.from_dict(json.loads(text), **meta)

    def clone(self, new_name: str) -> TemplateConfig:
        twin = TemplateConfig.from_dict(self.to_dict())
        twin.name = new_name
        return twin


def _c(header: str, source: str | None = None, default: str = "", max_length: int | None = None) -> ColumnSpec:
    return ColumnSpec(header, source, default, max_length)


_XERO_TAX = {
    "purchase:T0": "Zero Rated Expenses", "purchase:T1": "20% (VAT on Expenses)",
    "purchase:T2": "Exempt Expenses", "purchase:T5": "5% (VAT on Expenses)",
    "purchase:T9": "No VAT",
    "sales:T0": "Zero Rated Income", "sales:T1": "20% (VAT on Income)",
    "sales:T2": "Exempt Income", "sales:T5": "5% (VAT on Income)", "sales:T9": "No VAT",
}
_QBO_TAX = {"T0": "0.0% Z", "T1": "20.0% S", "T2": "Exempt", "T5": "5.0% R", "T9": "No VAT"}
_FREEAGENT_TAX = {"T0": "0", "T1": "20", "T2": "0", "T5": "5", "T9": "0"}

# Bump when a preset changes so existing databases pick up the new definition.
BUILTIN_TEMPLATE_VERSION = 1
_PRESET_NOTE = ("Built-in preset. Column names follow the vendor's import template at the time of "
                "writing — duplicate this template if your version of the software differs.")

BUILTIN_TEMPLATES: list[TemplateConfig] = [
    TemplateConfig(
        name="Sage 50 – Audit Trail Import",
        software="Sage 50",
        date_format="DD/MM/YYYY",
        encoding="cp1252",
        applies_to=["PI", "PC", "SI", "SC"],
        amount_sign="positive",
        notes=_PRESET_NOTE + " Import via File > Import > Audit Trail Transactions. "
              "Sage expects positive amounts; PC/SC types post credits.",
        columns=[
            _c("Type", F.TYPE),
            _c("Account Reference", F.ACCOUNT_CODE),
            _c("Nominal A/C Ref", F.NOMINAL_CODE),
            _c("Department Code", None, ""),
            _c("Date", F.INVOICE_DATE),
            _c("Reference", F.INVOICE_NUMBER, max_length=30),
            _c("Details", F.DESCRIPTION, max_length=60),
            _c("Net Amount", F.NET),
            _c("Tax Code", F.VAT_CODE),
            _c("Tax Amount", F.VAT),
            _c("Exchange Rate", F.EXCHANGE_RATE),
            _c("Extra Reference", F.SERVICE_PERIOD, max_length=30),
            _c("User Name", None, "LocalLedger"),
        ],
    ),
    TemplateConfig(
        name="Xero – Bills (Purchases)",
        software="Xero",
        date_format="DD/MM/YYYY",
        applies_to=["PI", "PC"],
        amount_sign="negative_for_credits",
        value_maps={F.VAT_CODE: dict(_XERO_TAX)},
        notes=_PRESET_NOTE + " Import via Business > Bills to pay > Import. Choose "
              "'tax exclusive' amounts. Negative lines import as credit notes. "
              "Date format follows your Xero organisation's region.",
        columns=[
            _c("*ContactName", F.COMPANY_NAME),
            _c("EmailAddress"),
            _c("*InvoiceNumber", F.INVOICE_NUMBER),
            _c("*InvoiceDate", F.INVOICE_DATE),
            _c("*DueDate", F.INVOICE_DATE),
            _c("Description", F.DESCRIPTION, max_length=4000),
            _c("*Quantity", None, "1"),
            _c("*UnitAmount", F.NET),
            _c("*AccountCode", F.NOMINAL_CODE),
            _c("*TaxType", F.VAT_CODE),
            _c("TaxAmount", F.VAT),
            _c("Currency", F.CURRENCY),
        ],
    ),
    TemplateConfig(
        name="Xero – Sales Invoices",
        software="Xero",
        date_format="DD/MM/YYYY",
        applies_to=["SI", "SC"],
        amount_sign="negative_for_credits",
        value_maps={F.VAT_CODE: dict(_XERO_TAX)},
        notes=_PRESET_NOTE + " Import via Business > Invoices > Import. Negative lines "
              "import as credit notes.",
        columns=[
            _c("*ContactName", F.COMPANY_NAME),
            _c("EmailAddress"),
            _c("*InvoiceNumber", F.INVOICE_NUMBER),
            _c("Reference", F.SERVICE_PERIOD),
            _c("*InvoiceDate", F.INVOICE_DATE),
            _c("*DueDate", F.INVOICE_DATE),
            _c("*Description", F.DESCRIPTION, max_length=4000),
            _c("*Quantity", None, "1"),
            _c("*UnitAmount", F.NET),
            _c("*AccountCode", F.NOMINAL_CODE),
            _c("*TaxType", F.VAT_CODE),
            _c("TaxAmount", F.VAT),
            _c("Currency", F.CURRENCY),
        ],
    ),
    TemplateConfig(
        name="QuickBooks Online – Bills",
        software="QuickBooks",
        date_format="DD/MM/YYYY",
        applies_to=["PI"],
        value_maps={F.VAT_CODE: dict(_QBO_TAX)},
        notes=_PRESET_NOTE + " QuickBooks bill import does not accept supplier credits; "
              "enter PC rows as Supplier Credits manually. US companies: switch the date "
              "format to MM/DD/YYYY.",
        columns=[
            _c("Bill No", F.INVOICE_NUMBER),
            _c("Supplier", F.COMPANY_NAME),
            _c("Bill Date", F.INVOICE_DATE),
            _c("Due Date", F.INVOICE_DATE),
            _c("Memo", F.SERVICE_PERIOD),
            _c("Account", F.NOMINAL_CODE),
            _c("Line Description", F.DESCRIPTION, max_length=4000),
            _c("Line Amount", F.NET),
            _c("Line Tax Code", F.VAT_CODE),
            _c("Line Tax Amount", F.VAT),
        ],
    ),
    TemplateConfig(
        name="QuickBooks Online – Invoices",
        software="QuickBooks",
        date_format="DD/MM/YYYY",
        applies_to=["SI"],
        value_maps={F.VAT_CODE: dict(_QBO_TAX)},
        notes=_PRESET_NOTE + " Credit notes (SC) must be raised in QuickBooks directly.",
        columns=[
            _c("InvoiceNo", F.INVOICE_NUMBER),
            _c("Customer", F.COMPANY_NAME),
            _c("InvoiceDate", F.INVOICE_DATE),
            _c("DueDate", F.INVOICE_DATE),
            _c("Item(Product/Service)", F.NOMINAL_CODE),
            _c("ItemDescription", F.DESCRIPTION, max_length=4000),
            _c("ItemQuantity", None, "1"),
            _c("ItemRate", F.NET),
            _c("ItemAmount", F.NET),
            _c("ItemTaxCode", F.VAT_CODE),
            _c("ItemTaxAmount", F.VAT),
            _c("Service Date", F.SERVICE_START),
        ],
    ),
    TemplateConfig(
        name="FreeAgent – Bills",
        software="FreeAgent",
        date_format="YYYY-MM-DD",
        applies_to=["PI", "PC"],
        amount_sign="negative_for_credits",
        value_maps={F.VAT_CODE: dict(_FREEAGENT_TAX)},
        notes=_PRESET_NOTE + " FreeAgent uses ISO dates and gross ('Total Value') amounts; "
              "map your nominal codes to FreeAgent category codes.",
        columns=[
            _c("Contact", F.COMPANY_NAME),
            _c("Reference", F.INVOICE_NUMBER),
            _c("Dated On", F.INVOICE_DATE),
            _c("Due On", F.INVOICE_DATE),
            _c("Category", F.NOMINAL_CODE),
            _c("Comments", F.DESCRIPTION),
            _c("Total Value", F.GROSS),
            _c("Sales Tax Rate", F.VAT_CODE),
            _c("Sales Tax Value", F.VAT),
            _c("Currency", F.CURRENCY),
        ],
    ),
    TemplateConfig(
        name="LocalLedger – Full Detail",
        software="Generic",
        date_format="YYYY-MM-DD",
        encoding="utf-8-sig",  # BOM so Excel opens UTF-8 correctly
        service_period_style="range",
        notes="Every internal field, ISO dates. Useful for review or as a base for custom templates.",
        columns=[_c(name, name) for name in (
            F.TYPE, F.ACCOUNT_CODE, F.COMPANY_NAME, F.NOMINAL_CODE, F.INVOICE_NUMBER,
            F.INVOICE_DATE, F.SERVICE_PERIOD, F.DESCRIPTION, F.NET, F.VAT_CODE, F.VAT,
            F.GROSS, F.CURRENCY, F.EXCHANGE_RATE, F.BASE_GROSS, F.SOURCE_FILE,
        )],
    ),
]
for _tpl in BUILTIN_TEMPLATES:
    _tpl.is_builtin = True
    _tpl.builtin_version = BUILTIN_TEMPLATE_VERSION


def to_base(amount: Decimal | None, currency: str, rate: Decimal | None, base_currency: str) -> Decimal | None:
    """Convert a document-currency amount to base currency.

    ``rate`` = units of document currency per 1 base unit, so base = amount / rate.
    Without a usable rate the amount is returned unchanged (flagged by the validator).
    """
    if amount is None:
        return None
    if (currency or base_currency).upper() == base_currency.upper() or not rate or rate <= 0:
        return amount
    return (amount / rate).quantize(CENT, rounding=ROUND_HALF_UP)


def service_month_of(row: StagingRow) -> date | None:
    """First day of the accrual month: majority month of the service period,
    else the period start, else the invoice date."""
    if row.service_period_start and row.service_period_end:
        return majority_month(row.service_period_start, row.service_period_end)
    anchor = row.service_period_start or row.service_period_end or row.invoice_date
    return anchor.replace(day=1) if anchor else None


def row_field_values(
    row: StagingRow,
    *,
    date_fmt: str = "%d/%m/%Y",
    base_currency: str = "GBP",
    service_period_style: str = "month",
) -> dict[str, Any]:
    """Internal field name -> value (dates already formatted, amounts as Decimal)."""

    def fmt(d: date | None) -> str:
        return d.strftime(date_fmt) if d else ""

    svc_month = service_month_of(row)
    if service_period_style == "range" and row.service_period_start and row.service_period_end:
        service_period = f"{fmt(row.service_period_start)} - {fmt(row.service_period_end)}"
    else:
        service_period = month_label(svc_month)

    rate = row.exchange_rate
    if rate is None and row.currency == base_currency.upper():
        rate = Decimal(1)

    return {
        F.TYPE: row.doc_type.value,
        F.ACCOUNT_CODE: row.account_code,
        F.NOMINAL_CODE: row.nominal_code,
        F.INVOICE_NUMBER: row.invoice_number,
        F.INVOICE_DATE: fmt(row.invoice_date),
        F.SERVICE_PERIOD: service_period,
        F.DESCRIPTION: row.description or row.raw_description,
        F.NET: row.net,
        F.VAT_CODE: row.vat_code,
        F.VAT: row.vat,
        F.GROSS: row.gross,
        F.CURRENCY: row.currency,
        F.EXCHANGE_RATE: rate,
        F.COMPANY_NAME: row.entity_name or row.company_name or row.account_code,
        F.SERVICE_MONTH: month_label(svc_month),
        F.SERVICE_START: fmt(row.service_period_start),
        F.SERVICE_END: fmt(row.service_period_end),
        F.RAW_DESCRIPTION: row.raw_description,
        F.LINE_NUMBER: str(row.line_no),
        F.SOURCE_FILE: row.source_file,
        F.BASE_NET: to_base(row.net, row.currency, row.exchange_rate, base_currency),
        F.BASE_VAT: to_base(row.vat, row.currency, row.exchange_rate, base_currency),
        F.BASE_GROSS: to_base(row.gross, row.currency, row.exchange_rate, base_currency),
    }


def lookup_value_map(mapping: Mapping[str, str] | None, value: str, doc_type: DocType) -> str:
    """Apply a value map; most specific key wins: 'PI:T1' > 'purchase:T1' > 'T1'."""
    if not mapping or value == "":
        return value
    for key in (f"{doc_type.value}:{value}", f"{doc_type.ledger}:{value}", value):
        for candidate in (key, key.upper(), key.lower()):
            if candidate in mapping:
                return mapping[candidate]
    return value


@dataclass
class MappingResult:
    template: TemplateConfig
    headers: list[str]
    kinds: list[str]                 # "amount" | "rate" | "text" per column
    records: list[list[Any]]         # typed values: Decimal for amounts/rates, str otherwise
    row_ids: list[str]               # staging row id per record
    excluded: list[tuple[str, str]] = field(default_factory=list)   # (row_id, reason)
    warnings: list[str] = field(default_factory=list)

    def _cell_text(self, value: Any, kind: str) -> str:
        if value is None:
            return ""
        if isinstance(value, Decimal):
            if kind == "rate":
                return _format_rate(value)
            return f"{value:.{self.template.decimal_places}f}"
        return str(value)

    def text_records(self) -> list[list[str]]:
        return [[self._cell_text(v, k) for v, k in zip(rec, self.kinds)] for rec in self.records]

    def to_dataframe(self):
        import pandas as pd

        return pd.DataFrame(self.text_records(), columns=self.headers)

    def to_csv_bytes(self) -> bytes:
        buffer = io.StringIO(newline="")
        writer = csv.writer(buffer, delimiter=self.template.delimiter,
                            quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
        if self.template.include_header:
            writer.writerow(self.headers)
        writer.writerows(self.text_records())
        return buffer.getvalue().encode(self.template.encoding, errors="replace")

    def to_excel_bytes(self) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Upload"
        start = 1
        if self.template.include_header:
            ws.append(self.headers)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            ws.freeze_panes = "A2"
            start = 2
        number_format = "0." + "0" * self.template.decimal_places if self.template.decimal_places else "0"
        for r_idx, record in enumerate(self.records, start=start):
            for c_idx, (value, kind) in enumerate(zip(record, self.kinds), start=1):
                if isinstance(value, Decimal):
                    cell = ws.cell(row=r_idx, column=c_idx, value=float(value))
                    cell.number_format = number_format if kind == "amount" else "0.000000"
                else:
                    ws.cell(row=r_idx, column=c_idx, value="" if value is None else str(value))
        for c_idx, header in enumerate(self.headers, start=1):
            longest = max([len(header)] + [len(self._cell_text(rec[c_idx - 1], self.kinds[c_idx - 1]))
                                           for rec in self.records[:200]])
            ws.column_dimensions[get_column_letter(c_idx)].width = min(max(10, longest + 2), 60)
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    def to_bytes(self) -> bytes:
        return self.to_excel_bytes() if self.template.file_format == "xlsx" else self.to_csv_bytes()


def _format_rate(value: Decimal) -> str:
    return format(value.normalize(), "f")


def map_rows(
    rows: Iterable[StagingRow],
    template: TemplateConfig,
    *,
    base_currency: str = "GBP",
) -> MappingResult:
    """Transform staging rows into the template's output layout."""
    errors = template.validate()
    if errors:
        raise ValueError("Invalid template: " + "; ".join(errors))

    quantum = Decimal(1).scaleb(-template.decimal_places)
    kinds = [
        "amount" if c.source in AMOUNT_FIELDS else "rate" if c.source == F.EXCHANGE_RATE else "text"
        for c in template.columns
    ]
    allowed = set(template.applies_to)
    result = MappingResult(template, [c.header for c in template.columns], kinds, [], [])
    truncated: dict[str, int] = {}

    for row in rows:
        if row.doc_type.value not in allowed:
            result.excluded.append(
                (row.row_id, f"{row.doc_type.label} rows are not supported by '{template.name}'."))
            continue
        values = row_field_values(row, date_fmt=template.strftime, base_currency=base_currency,
                                  service_period_style=template.service_period_style)
        sign = -1 if (template.amount_sign == "negative_for_credits" and row.doc_type.is_credit) else 1
        record: list[Any] = []
        for col, kind in zip(template.columns, kinds):
            value: Any
            if col.source is None:
                value = col.default
            elif kind == "amount":
                raw = values.get(col.source)
                value = None if raw is None else (raw * sign).quantize(quantum, rounding=ROUND_HALF_UP)
            elif kind == "rate":
                value = values.get(col.source)
            else:
                raw = values.get(col.source)
                value = lookup_value_map(template.value_maps.get(col.source),
                                         "" if raw is None else str(raw), row.doc_type)
            if value is None or value == "":
                value = col.default
            if isinstance(value, str) and col.max_length and len(value) > col.max_length:
                value = value[: col.max_length].rstrip()
                truncated[col.header] = truncated.get(col.header, 0) + 1
            record.append(value)
        result.records.append(record)
        result.row_ids.append(row.row_id)

    for header, count in truncated.items():
        result.warnings.append(f"{count} value(s) in '{header}' were truncated to fit the column limit.")
    return result


# =============================================================================
# 2. SERVICE PERIOD DETECTION
# =============================================================================
_MONTHS = (r"jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|june?|july?|aug(?:ust)?|"
           r"sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?")
_DATE = (rf"(?:\d{{4}}-\d{{1,2}}-\d{{1,2}}"
         rf"|\d{{1,2}}[/.\-]\d{{1,2}}[/.\-]\d{{2,4}}"
         rf"|\d{{1,2}}(?:st|nd|rd|th)?[\s\-](?:{_MONTHS})\.?,?[\s\-]\d{{2,4}}"
         rf"|(?:{_MONTHS})\.?\s+\d{{1,2}}(?:st|nd|rd|th)?,?\s+\d{{4}})")
DATE_PATTERN = _DATE   # public: reused by the extraction parsers
MONTHS_PATTERN = _MONTHS
_SEP = r"\s*(?:-|–|—|to|until|till|through|thru)\s*"
_KEYWORDS = (r"(?:(?:billing|service|invoice|invoicing|statement|accounting|subscription|rental|hire|"
             r"usage|charge|charging|licen[cs]e|contract|coverage|supply)\s+)?period"
             r"|period\s+(?:covered|of\s+service)|service\s+dates?|dates?\s+of\s+service"
             r"|billed\s+for|covering")
_RANGE_KEYWORD_RE = re.compile(
    rf"(?:{_KEYWORDS})\s*(?:\(.*?\))?\s*[:\-]?\s*(?:from\s+)?(?P<start>{_DATE}){_SEP}(?P<end>{_DATE})",
    re.I)
_FROM_TO_RE = re.compile(rf"\bfrom\s+(?P<start>{_DATE})\s*(?:to|until|till|through|-|–)\s*(?P<end>{_DATE})", re.I)
_MONTH_WORD_RE = re.compile(
    rf"(?:{_KEYWORDS}|service\s+month|billing\s+month|(?:for\s+the\s+)?month\s+of)\s*[:\-]?\s*"
    rf"(?P<month>{_MONTHS})\.?,?\s*(?P<year>\d{{4}})\b", re.I)
_MONTH_NUM_RE = re.compile(
    rf"(?:{_KEYWORDS}|service\s+month|billing\s+month)\s*[:\-]?\s*(?P<month>\d{{1,2}})[/\-.](?P<year>\d{{4}})\b",
    re.I)
_MONTH_INDEX = {m: i for i, m in enumerate(
    ("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"), start=1)}
_MAX_PERIOD_DAYS = 400


@dataclass(frozen=True)
class ServicePeriod:
    start: date | None
    end: date | None
    source: str          # "range" | "month" | "invoice_date" | "none"
    matched_text: str = ""

    @property
    def month(self) -> date | None:
        if self.start and self.end:
            return majority_month(self.start, self.end)
        return self.start.replace(day=1) if self.start else None

    @property
    def month_label(self) -> str:
        return month_label(self.month)

    @property
    def period_label(self) -> str:
        if not (self.start and self.end):
            return self.month_label
        first, last = month_label(self.start), month_label(self.end)
        return first if first == last else f"{first} - {last}"


def _end_of_month(d: date) -> date:
    nxt = date(d.year + (d.month == 12), d.month % 12 + 1, 1)
    return nxt - timedelta(days=1)


def majority_month(start: date, end: date) -> date:
    """First day of the calendar month holding most days of [start, end].
    Ties go to the earlier month."""
    if end < start:
        start, end = end, start
    best: tuple[int, date] | None = None
    cursor = start
    while cursor <= end:
        segment_end = min(end, _end_of_month(cursor))
        days = (segment_end - cursor).days + 1
        month_start = cursor.replace(day=1)
        if best is None or days > best[0]:
            best = (days, month_start)
        cursor = segment_end + timedelta(days=1)
    assert best is not None
    return best[1]


def detect_service_period(
    text: str | None,
    invoice_date: date | None = None,
    *,
    dayfirst: bool = True,
) -> ServicePeriod:
    """Find a billing/service period in document text.

    Order: keyword + date range, 'from X to Y', keyword + month name,
    keyword + MM/YYYY, then falls back to the invoice date's month.
    """
    if text:
        for regex in (_RANGE_KEYWORD_RE, _FROM_TO_RE):
            for match in regex.finditer(text):
                start = parse_date(match["start"], dayfirst=dayfirst)
                end = parse_date(match["end"], dayfirst=dayfirst)
                if start and end and start <= end and (end - start).days <= _MAX_PERIOD_DAYS:
                    return ServicePeriod(start, end, "range", match.group(0).strip())
        for match in _MONTH_WORD_RE.finditer(text):
            month = _MONTH_INDEX.get(match["month"][:3].lower())
            if month:
                start = date(int(match["year"]), month, 1)
                return ServicePeriod(start, _end_of_month(start), "month", match.group(0).strip())
        for match in _MONTH_NUM_RE.finditer(text):
            month, year = int(match["month"]), int(match["year"])
            if 1 <= month <= 12 and 1990 <= year <= 2100:
                start = date(year, month, 1)
                return ServicePeriod(start, _end_of_month(start), "month", match.group(0).strip())
    if invoice_date:
        start = invoice_date.replace(day=1)
        return ServicePeriod(start, _end_of_month(start), "invoice_date")
    return ServicePeriod(None, None, "none")


# =============================================================================
# 3. DESCRIPTION BUILDER
# =============================================================================
DESCRIPTION_TOKENS: dict[str, str] = {
    "Raw_Description": "Description text extracted from the document",
    "Company_Name": "Supplier/client name (master-list name when matched)",
    "Service_Month": "Accrual month, e.g. 'Oct 2026' (falls back to invoice month)",
    "Service_Period": "Period, e.g. 'Sep 2026' or 'Jul 2026 - Sep 2026'",
    "Invoice_Month": "Month of the invoice date, e.g. 'Oct 2026'",
    "Invoice_Number": "Invoice / credit note number",
    "Invoice_Date": "Invoice date (DD/MM/YYYY)",
    "Account_Code": "Supplier/client account code",
    "Nominal_Code": "Nominal (GL) code",
    "Type": "PI, PC, SI or SC",
    "Currency": "Document currency code",
}
_TOKEN_RE = re.compile(r"\{([A-Za-z_]+)\}")
_TOKEN_LOOKUP = {k.lower(): k for k in DESCRIPTION_TOKENS}
_SEPARATORS = "-–—|:,;."
_CASES = ("none", "upper", "lower", "title", "sentence")


@dataclass
class DescriptionRule:
    """A token template for the Description column.

    The rendered value is ``prefix + template + suffix``. A rule applies when
    the row's type is in ``doc_types`` and its account in ``account_codes``
    (empty lists mean "any"). The most specific matching rule wins.
    """

    name: str = "Default"
    template: str = "{Raw_Description}"
    prefix: str = ""
    suffix: str = ""
    doc_types: list[str] = field(default_factory=list)
    account_codes: list[str] = field(default_factory=list)
    max_length: int | None = None
    case: str = "none"
    enabled: bool = True

    def matches(self, row: StagingRow) -> bool:
        if not self.enabled:
            return False
        if self.doc_types and row.doc_type.value not in {t.upper() for t in self.doc_types}:
            return False
        if self.account_codes and row.account_code.upper() not in {a.upper() for a in self.account_codes}:
            return False
        return True

    @property
    def specificity(self) -> int:
        return (2 if self.account_codes else 0) + (1 if self.doc_types else 0)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> DescriptionRule:
        max_length = data.get("max_length")
        return cls(
            name=str(data.get("name", "Rule")),
            template=str(data.get("template", "{Raw_Description}")),
            prefix=str(data.get("prefix", "")),
            suffix=str(data.get("suffix", "")),
            doc_types=[str(t).upper() for t in data.get("doc_types") or []],
            account_codes=[str(a) for a in data.get("account_codes") or []],
            max_length=int(max_length) if max_length not in (None, "", 0) else None,
            case=data.get("case", "none") if data.get("case") in _CASES else "none",
            enabled=bool(data.get("enabled", True)),
        )


DEFAULT_DESCRIPTION_RULES: list[DescriptionRule] = [
    DescriptionRule(
        name="Default",
        prefix="{Service_Month} - ",
        template="{Raw_Description}",
        suffix=". {Company_Name}",
        max_length=60,
    ),
]

SAMPLE_ROW = StagingRow(
    doc_type=DocType.PI,
    account_code="ACME001",
    nominal_code="7100",
    invoice_number="INV-10452",
    invoice_date=date(2026, 10, 3),
    service_period_start=date(2026, 9, 1),
    service_period_end=date(2026, 9, 30),
    raw_description="Office rent",
    company_name="ACME Property Services Ltd",
    entity_name="ACME Property Services Ltd",
    net=Decimal("1000.00"), vat_code="T1", vat=Decimal("200.00"), gross=Decimal("1200.00"),
)


def description_context(row: StagingRow, *, date_fmt: str = "%d/%m/%Y") -> dict[str, str]:
    period = ServicePeriod(row.service_period_start, row.service_period_end,
                           "range" if row.service_period_start else "none")
    svc = service_month_of(row)
    return {
        "Raw_Description": row.raw_description.strip(),
        "Company_Name": (row.entity_name or row.company_name).strip(),
        "Service_Month": month_label(svc),
        "Service_Period": period.period_label if row.service_period_start else month_label(svc),
        "Invoice_Month": month_label(row.invoice_date),
        "Invoice_Number": row.invoice_number.strip(),
        "Invoice_Date": row.invoice_date.strftime(date_fmt) if row.invoice_date else "",
        "Account_Code": row.account_code.strip(),
        "Nominal_Code": row.nominal_code.strip(),
        "Type": row.doc_type.value,
        "Currency": row.currency,
    }


def validate_description_template(text: str) -> list[str]:
    """Return human-readable problems (unknown tokens, unbalanced braces)."""
    problems = []
    unknown = sorted({t for t in _TOKEN_RE.findall(text) if t.lower() not in _TOKEN_LOOKUP})
    if unknown:
        problems.append("Unknown token(s): " + ", ".join("{" + t + "}" for t in unknown)
                        + ". Available: " + ", ".join("{" + t + "}" for t in DESCRIPTION_TOKENS))
    stripped = _TOKEN_RE.sub("", text)
    if "{" in stripped or "}" in stripped:
        problems.append("Unbalanced '{' or '}' — tokens look like {Service_Month}.")
    return problems


def _substitute(text: str, context: Mapping[str, str]) -> str:
    def repl(match: re.Match[str]) -> str:
        key = _TOKEN_LOOKUP.get(match.group(1).lower())
        return context.get(key, "") if key else match.group(0)

    return _TOKEN_RE.sub(repl, text)


def _tidy(text: str) -> str:
    """Remove separators orphaned by empty tokens and normalise whitespace."""
    text = re.sub(r"\s+", " ", text).strip()
    # "Sep 2026 - . ACME" -> "Sep 2026 - ACME"
    sep = re.escape(_SEPARATORS)
    text = re.sub(rf"(\s[{sep}])(?:\s*[{sep}](?=\s|$))+", r"\1", text)
    text = re.sub(r"\s+([.,;:])", r"\1", text)
    text = text.strip(" " + _SEPARATORS)
    # drop consecutive duplicate " - " segments ("Sep 2026 - Sep 2026 - Rent")
    parts = text.split(" - ")
    deduped = [p for i, p in enumerate(parts) if i == 0 or p.strip().lower() != parts[i - 1].strip().lower()]
    return " - ".join(deduped).strip()


def _apply_case(text: str, case: str) -> str:
    if case == "upper":
        return text.upper()
    if case == "lower":
        return text.lower()
    if case == "title":
        return text.title()
    if case == "sentence":
        return text[:1].upper() + text[1:]
    return text


def _truncate(text: str, limit: int | None) -> str:
    if not limit or len(text) <= limit:
        return text
    cut = text[:limit]
    space = cut.rfind(" ")
    if space >= int(limit * 0.6):
        cut = cut[:space]
    return cut.rstrip(" " + _SEPARATORS)


def render_description(rule: DescriptionRule, context: Mapping[str, str]) -> str:
    body = _tidy(_substitute(rule.template, context))
    prefix = _tidy(_substitute(rule.prefix, context)) if rule.prefix else ""
    suffix = _tidy(_substitute(rule.suffix, context)) if rule.suffix else ""
    # Don't repeat a prefix/suffix the document text already contains.
    if prefix and body.lower().startswith(prefix.lower()):
        prefix = ""
    if suffix and body.lower().endswith(suffix.lower()):
        suffix = ""
    prefix_raw = _substitute(rule.prefix, context) if prefix else ""
    suffix_raw = _substitute(rule.suffix, context) if suffix else ""
    body_raw = _substitute(rule.template, context)
    full = _apply_case(_tidy(prefix_raw + body_raw + suffix_raw), rule.case)
    limit = rule.max_length
    if not limit or len(full) <= limit:
        return full
    # Over the limit: shorten the body first so the prefix (month) and suffix
    # (company) survive, e.g. "Sep 2026 - Office rent. ACME Property Services Ltd".
    overflow = len(full) - limit
    body = _tidy(body_raw)
    if len(body) - overflow >= 8:
        shortened = _truncate(body, len(body) - overflow)
        full = _apply_case(_tidy(prefix_raw + shortened + suffix_raw), rule.case)
    return _truncate(full, limit)


def select_rule(row: StagingRow, rules: Sequence[DescriptionRule]) -> DescriptionRule | None:
    candidates = [(idx, r) for idx, r in enumerate(rules) if r.matches(row)]
    if not candidates:
        return None
    return sorted(candidates, key=lambda pair: (-pair[1].specificity, pair[0]))[0][1]


def build_description(row: StagingRow, rules: Sequence[DescriptionRule]) -> str:
    """Render the Description for ``row`` with the best matching rule.
    Without a matching rule the raw description is returned unchanged."""
    rule = select_rule(row, rules)
    if rule is None:
        return row.raw_description.strip()
    return render_description(rule, description_context(row))


def preview_description(rule: DescriptionRule, sample: StagingRow | None = None) -> str:
    """Live preview for the rule editor."""
    return render_description(rule, description_context(sample or SAMPLE_ROW))


def load_description_rules(data: Any) -> list[DescriptionRule]:
    """Accept a JSON string or list of dicts (as stored in settings)."""
    if isinstance(data, str):
        data = json.loads(data or "[]")
    return [r if isinstance(r, DescriptionRule) else DescriptionRule.from_dict(r) for r in data or []]


# =============================================================================
# 4. CREDIT NOTE DETECTION
# =============================================================================
_CREDIT_STRONG = re.compile(
    r"\bcredit[\s\-]*note\b|\bcredit[\s\-]*memo(?:randum)?\b|\bcredit\s+invoice\b|"
    r"\bcredit\s+advice\b|\bgutschrift\b|\bavoir\b|\bnota\s+de\s+cr[eé]dito\b|\bcreditnota\b",
    re.I)
_CREDIT_WEAK = re.compile(
    r"\brefund(?:ed)?\b|\bcredited\s+to\s+your\s+account\b|\breversal\b|\bcancell?ation\s+invoice\b|"
    r"\bamount\s+credited\b|\breturns?\s+credit\b",
    re.I)
_CREDIT_FALSE_POSITIVES = re.compile(
    r"credit\s*card|credit\s*terms?|credit\s*limit|credit\s*control|credit\s*check|"
    r"credit\s*account|credit\s*rating|letter\s+of\s+credit|refund\s*policy|no\s+refunds?|"
    r"non[\s\-]?refundable|refunds?\s+(?:will|are|may|can)\s+(?:only\s+)?be|"
    r"any\s+credit\s+notes?\s+will",
    re.I)
_INVOICE_TITLE = re.compile(r"\b(?:tax\s+)?invoice\b", re.I)
_AMT = r"\d[\d,]*\.\d{2}"
_NEGATIVE_TOTAL = re.compile(
    r"\b(?:grand\s+total|total(?:\s+(?:due|amount|payable|credit|gbp|eur|usd|inc\.?\s*vat))?|"
    r"amount\s+(?:due|payable|credited)|balance(?:\s+due)?)\b[^\n\d\-(]{0,30}"
    rf"(?:-\s?[£$€]?\s?{_AMT}|\(\s?[£$€]?\s?{_AMT}\s?\)|[£$€]?\s?{_AMT}\s?CR\b)",
    re.I)


@dataclass(frozen=True)
class CreditNoteDetection:
    is_credit: bool
    score: float               # 0-1
    reasons: tuple[str, ...]


def detect_credit_note(
    text: str | None,
    *,
    net: Decimal | None = None,
    gross: Decimal | None = None,
    header_lines: int = 20,
    threshold: float = 0.6,
) -> CreditNoteDetection:
    """Decide whether a document is a credit note.

    Signals: strong keywords ("Credit Note", "Credit Memo") — weighted higher
    near the top of the page; weak keywords ("Refund"); a negative total in
    the text; negative extracted amounts. Boilerplate such as "credit card"
    or "refund policy" is ignored.
    """
    score = 0.0
    reasons: list[str] = []
    cleaned = _CREDIT_FALSE_POSITIVES.sub(" ", text or "")
    lines = [ln for ln in cleaned.splitlines() if ln.strip()]
    header = "\n".join(lines[:header_lines])

    strong_header = _CREDIT_STRONG.search(header)
    strong_any = strong_header or _CREDIT_STRONG.search(cleaned)
    if strong_header:
        score += 0.9
        reasons.append(f"'{strong_header.group(0)}' in document header")
    elif strong_any:
        score += 0.5
        reasons.append(f"'{strong_any.group(0)}' in document body")

    weak_hits = {m.group(0).lower() for m in _CREDIT_WEAK.finditer(cleaned)}
    if weak_hits:
        score += min(0.5, 0.25 * len(weak_hits))
        reasons.append("keyword(s): " + ", ".join(sorted(weak_hits)))

    neg = _NEGATIVE_TOTAL.search(cleaned)
    if neg:
        score += 0.6
        reasons.append(f"negative total '{neg.group(0).strip()}'")

    reference_amount = gross if gross is not None else net
    if reference_amount is not None and reference_amount < 0:
        score += 0.7
        reasons.append("negative extracted total")

    if not strong_header and _INVOICE_TITLE.search(header):
        score -= 0.2

    score = max(0.0, min(1.0, score))
    return CreditNoteDetection(score >= threshold, round(score, 2), tuple(reasons))


def apply_credit_detection(row: StagingRow, detection: CreditNoteDetection | None = None) -> list[str]:
    """Flip PI->PC / SI->SC when a credit note is detected, and make amounts
    positive (the sign lives in the type). Returns a list of changes."""
    changes: list[str] = []
    negative = any(v is not None and v < 0 for v in (row.net, row.vat, row.gross, row.doc_gross))
    if (detection and detection.is_credit) or negative:
        if not row.doc_type.is_credit:
            old = row.doc_type
            row.doc_type = row.doc_type.as_credit()
            why = "; ".join(detection.reasons) if detection and detection.is_credit else "negative amounts"
            changes.append(f"Type changed {old.value} -> {row.doc_type.value} ({why})")
    if negative:
        for name in ("net", "vat", "gross", "doc_gross"):
            value = getattr(row, name)
            if value is not None and value < 0:
                setattr(row, name, -value)
        changes.append("Negative amounts converted to positive (credit type carries the sign)")
    return changes


# =============================================================================
# 5. ARCHIVE NAMING
# =============================================================================
_FILENAME_TOKEN_RE = re.compile(r"\[([A-Za-z]+)\]")
_INVALID_FILENAME_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_WINDOWS_RESERVED = {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)),
                     *(f"LPT{i}" for i in range(1, 10))}
FILENAME_TOKENS = ("AccountCode", "InvoiceDate", "InvoiceNumber", "Type", "CompanyName",
                   "Gross", "Currency", "ServiceMonth", "OriginalName")


def _sanitize_component(text: str) -> str:
    text = _INVALID_FILENAME_CHARS.sub("-", text)
    text = re.sub(r"\s+", " ", text).strip(" .")
    return text


def render_filename(pattern: str, row: StagingRow, *, date_fmt: str = "%Y-%m-%d",
                    max_stem: int = 150) -> str:
    """Render e.g. ``[AccountCode]_[InvoiceDate]_[InvoiceNumber].pdf`` into a
    safe file name for Windows, macOS and Linux."""
    svc = service_month_of(row)
    original = re.sub(r"\.pdf$", "", row.source_file or "", flags=re.I)
    tokens = {
        "accountcode": row.account_code,
        "invoicedate": row.invoice_date.strftime(date_fmt) if row.invoice_date else "",
        "invoicenumber": row.invoice_number,
        "type": row.doc_type.value,
        "companyname": row.entity_name or row.company_name,
        "gross": f"{row.doc_gross or row.gross or 0:.2f}",
        "currency": row.currency,
        "servicemonth": svc.strftime("%Y-%m") if svc else "",
        "originalname": original,
    }

    def repl(match: re.Match[str]) -> str:
        key = match.group(1).lower()
        return _sanitize_component(str(tokens[key])) if key in tokens else match.group(0)

    name = _FILENAME_TOKEN_RE.sub(repl, pattern or "[OriginalName].pdf")
    name = _INVALID_FILENAME_CHARS.sub("-", name)
    stem, dot, ext = name.rpartition(".")
    if not dot or len(ext) > 5 or " " in ext:
        stem, ext = name, "pdf"
    stem = re.sub(r"([_\-\s])\1+", r"\1", stem)
    stem = re.sub(r"[_\-\s]{2,}", "_", stem).strip(" _-.")
    stem = stem[:max_stem].rstrip(" _-.") or (original or "document")
    if stem.upper() in _WINDOWS_RESERVED:
        stem = f"_{stem}"
    return f"{stem}.{ext.lower()}"


def archive_subfolder(pattern: str, row: StagingRow, *, today: date | None = None) -> str:
    """Render the archive sub-folder, e.g. ``{YYYY}-{MM}`` -> ``2026-10``.
    Tokens: {YYYY} {MM} {Type} {Ledger} {AccountCode}. Uses the invoice date."""
    anchor = row.invoice_date or today or date.today()
    tokens = {
        "YYYY": f"{anchor.year:04d}", "MM": f"{anchor.month:02d}",
        "TYPE": row.doc_type.value, "LEDGER": "Sales" if row.doc_type.is_sales else "Purchases",
        "ACCOUNTCODE": row.account_code or "UNKNOWN",
    }
    rendered = re.sub(r"\{([A-Za-z]+)\}",
                      lambda m: tokens.get(m.group(1).upper(), m.group(0)), pattern or "{YYYY}-{MM}")
    parts = [_sanitize_component(p) for p in re.split(r"[\\/]+", rendered)]
    return "/".join(p for p in parts if p and p not in {".", ".."})


# =============================================================================
# 6. PIPELINE CONVENIENCE
# =============================================================================
def apply_rules(
    row: StagingRow,
    *,
    text: str | None = None,
    description_rules: Sequence[DescriptionRule] | None = None,
    dayfirst: bool = True,
    detect_credit: bool = True,
) -> list[str]:
    """Run credit detection, service-period detection and the description
    builder on a freshly extracted row. Returns a human-readable change log."""
    changes: list[str] = []
    source_text = text if text is not None else row.raw_text
    if detect_credit:
        detection = detect_credit_note(source_text, net=row.net, gross=row.gross)
        changes += apply_credit_detection(row, detection)
    if not row.service_period_start:
        period = detect_service_period(source_text, row.invoice_date, dayfirst=dayfirst)
        if period.source in ("range", "month"):
            row.service_period_start, row.service_period_end = period.start, period.end
            changes.append(f"Service period {period.period_label} detected from '{period.matched_text}'")
    if description_rules:
        new_description = build_description(row, description_rules)
        if new_description != row.description:
            row.description = new_description
            changes.append(f"Description set to '{new_description}'")
    return changes
